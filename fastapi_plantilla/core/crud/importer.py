import contextlib
import csv
import io
import json
from collections.abc import Callable
from dataclasses import dataclass
from typing import Any, Final

from fastapi import HTTPException, status
from loguru import logger
from pydantic import BaseModel, ValidationError
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from fastapi_plantilla.core.crud.schema import (
    ImportFormat,
    ImportJobPayload,
    ImportMode,
    ImportResult,
    RowError,
    ScopeContext,
    WriteOptions,
)
from fastapi_plantilla.modules.jobs.exceptions import JobCancelledError, JobError
from fastapi_plantilla.modules.jobs.registry import register_job
from fastapi_plantilla.modules.jobs.schema import JobContext

try:
    import openpyxl  # type: ignore[import-not-found]
except ImportError:
    openpyxl = None  # type: ignore[assignment]

__all__ = [
    "DEFAULT_MAX_IMPORT_FILE_BYTES",
    "DEFAULT_MAX_IMPORT_ROWS",
    "IMPORT_STORAGE_PREFIX",
    "MAX_EMBEDDED_ERRORS",
    "RESOURCE_IMPORT_REGISTRY",
    "SYSTEM_IMPORT_EXCLUDE_FIELDS",
    "ResourceImportDefinition",
    "generate_import_template",
    "handle_import_job",
    "parse_import_file",
    "register_import_resource",
]

DEFAULT_MAX_IMPORT_ROWS: Final[int] = 5000
DEFAULT_MAX_IMPORT_FILE_BYTES: Final[int] = 10 * 1024 * 1024  # 10 MB
MAX_EMBEDDED_ERRORS: Final[int] = 100
IMPORT_STORAGE_PREFIX: Final[str] = "imports"
SYSTEM_IMPORT_EXCLUDE_FIELDS: Final[frozenset[str]] = frozenset(
    {
        "id",
        "created_at",
        "updated_at",
        "deleted_at",
        "restored_at",
        "version",
        "created_by",
        "updated_by",
        "deleted_by",
        "restored_by",
        "owner_id",
        "team_id",
        "creator",
        "updater",
        "status",
    }
)


@dataclass(frozen=True)
class ResourceImportDefinition:
    """Metadata and service factory for importable domain resources."""

    resource_name: str
    schema_create: type[BaseModel]
    service_factory: Callable[[AsyncSession], Any]


RESOURCE_IMPORT_REGISTRY: dict[str, ResourceImportDefinition] = {}


def register_import_resource(
    resource_name: str,
    schema_create: type[BaseModel],
    service_factory: Callable[[AsyncSession], Any],
) -> None:
    """Register a domain resource for mass data import handling."""
    RESOURCE_IMPORT_REGISTRY[resource_name.lower()] = ResourceImportDefinition(
        resource_name=resource_name.lower(),
        schema_create=schema_create,
        service_factory=service_factory,
    )


def generate_import_template(
    schema: type[BaseModel],
    fmt: ImportFormat,
    resource_name: str,
) -> tuple[bytes, str, str]:
    """Generate blank import template (.xlsx or .csv) with schema field headers."""
    fields: list[tuple[str, bool, str]] = []
    for name, field_info in schema.model_fields.items():
        if name in SYSTEM_IMPORT_EXCLUDE_FIELDS:
            continue
        is_req = field_info.is_required()
        desc = field_info.description or ("Required" if is_req else "Optional")
        fields.append((name, is_req, desc))

    headers = [f[0] for f in fields]
    clean_resource = resource_name.lower().replace(" ", "_")

    if fmt == ImportFormat.CSV:
        out = io.StringIO()
        writer = csv.writer(out)
        writer.writerow(headers)
        return (
            out.getvalue().encode("utf-8-sig"),
            "text/csv; charset=utf-8",
            f"{clean_resource}_template.csv",
        )

    if openpyxl is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Excel format requires openpyxl. Use CSV instead.",
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{clean_resource[:25]} Import"
    ws.append(headers)
    out_b = io.BytesIO()
    wb.save(out_b)
    return (
        out_b.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        f"{clean_resource}_template.xlsx",
    )


def parse_import_file(content: bytes, filename: str) -> list[dict[str, Any]]:
    """Parse raw bytes into normalized list of row dictionaries."""
    lower_fn = filename.lower()
    if lower_fn.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        dialect: Any = csv.excel
        with io.StringIO(text) as s_io:
            sample = s_io.read(2048)
            s_io.seek(0)
            if sample:
                with contextlib.suppress(csv.Error):
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            reader = csv.DictReader(s_io, dialect=dialect)
            return [
                {
                    k.strip(): v.strip() if isinstance(v, str) else v
                    for k, v in row.items()
                    if k
                }
                for row in reader
                if any(v is not None and str(v).strip() != "" for v in row.values())
            ]

    if openpyxl is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Excel parsing requires openpyxl. Upload CSV instead.",
        )

    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return []
    headers = [str(c).strip() if c is not None else "" for c in header_row]
    parsed_rows: list[dict[str, Any]] = []
    for row in rows_iter:
        if not any(c is not None and str(c).strip() != "" for c in row):
            continue
        row_dict: dict[str, Any] = {}
        for h, val in zip(headers, row, strict=False):
            if h:
                row_dict[h] = val.strip() if isinstance(val, str) else val
        parsed_rows.append(row_dict)
    return parsed_rows


async def _validate_rows(
    raw_rows: list[dict[str, Any]],
    schema: type[BaseModel],
    ctx: JobContext[ImportJobPayload],
) -> tuple[list[tuple[int, BaseModel]], list[RowError]]:
    """Validate all rows against Pydantic schema with throttled progress updates."""
    total_rows = len(raw_rows)
    throttle_step = max(1, total_rows // 20)
    errors: list[RowError] = []
    valid_rows: list[tuple[int, BaseModel]] = []

    await ctx.update_progress(5, f"Validating {total_rows} rows...")
    for idx, row_dict in enumerate(raw_rows, start=1):
        if idx % 100 == 0 and await ctx.is_cancelled():
            raise JobCancelledError("Import job cancelled by user.")
        try:
            instance = schema.model_validate(row_dict)
            valid_rows.append((idx, instance))
        except ValidationError as exc:
            for err in exc.errors():
                field_name = str(err["loc"][0]) if err["loc"] else "general"
                errors.append(
                    RowError(
                        row=idx,
                        field=field_name,
                        message=err["msg"],
                        value=row_dict.get(field_name),
                    )
                )
        if idx % throttle_step == 0:
            pct = 5 + int(45 * (idx / max(1, total_rows)))
            await ctx.update_progress(pct, f"Validated {idx}/{total_rows} rows")

    return valid_rows, errors


async def _persist_rows(
    service: Any,
    valid_rows: list[tuple[int, BaseModel]],
    payload: ImportJobPayload,
    scope: ScopeContext | None,
    options: WriteOptions,
    session: AsyncSession,
    ctx: JobContext[ImportJobPayload],
    errors: list[RowError],
) -> int:
    """Persist validated rows according to ATOMIC or PARTIAL transactional mode."""
    if payload.mode == ImportMode.ATOMIC:
        if errors:
            return 0
        await ctx.update_progress(60, "Persisting records in database...")
        try:
            async with session.begin_nested():
                res = await service.bulk_create(
                    [item for _, item in valid_rows],
                    user_id=payload.user_id,
                    scope=scope,
                    options=options,
                )
            return res.count if hasattr(res, "count") else len(valid_rows)
        except (HTTPException, IntegrityError) as exc:
            logger.warning("Error persisting atomic import batch: {}", exc)
            detail = exc.detail if isinstance(exc, HTTPException) else str(exc)
            errors.append(RowError(row=0, field="database", message=str(detail)))
            return 0

    # PARTIAL mode
    await ctx.update_progress(60, "Persisting valid rows...")
    imported_count = 0
    for idx, inst in valid_rows:
        try:
            async with session.begin_nested():
                await service.create(
                    data=inst,
                    user_id=payload.user_id,
                    scope=scope,
                    options=options,
                )
            imported_count += 1
        except (HTTPException, IntegrityError) as exc:
            logger.warning("Error persisting row {}: {}", idx, exc)
            detail = exc.detail if isinstance(exc, HTTPException) else str(exc)
            errors.append(RowError(row=idx, field="database", message=str(detail)))
    return imported_count


@register_job("imports.validate", payload_model=ImportJobPayload)
async def handle_import_job(ctx: JobContext[ImportJobPayload]) -> dict[str, Any]:
    """Worker handling row-by-row Pydantic validation, persistence, and reporting."""
    from fastapi_plantilla.modules.storage.dependencies import (  # noqa: PLC0415
        get_storage_provider,
    )

    payload = ctx.payload
    storage = get_storage_provider()
    raw_content: bytes = b""
    try:
        raw_content = await storage.download(payload.storage_key)
    finally:
        with contextlib.suppress(Exception):
            await storage.delete(payload.storage_key)

    resource_def = RESOURCE_IMPORT_REGISTRY.get(payload.resource_name.lower())
    if not resource_def:
        raise JobError(
            f"Resource '{payload.resource_name}' is not registered for import."
        )

    raw_rows = parse_import_file(raw_content, payload.filename)
    total_rows = len(raw_rows)
    if total_rows == 0:
        raise JobError("File is empty or contains no valid data rows.")
    if total_rows > DEFAULT_MAX_IMPORT_ROWS:
        msg = (
            f"File contains {total_rows} rows, exceeding limit "
            f"of {DEFAULT_MAX_IMPORT_ROWS}."
        )
        raise JobError(msg)

    schema = resource_def.schema_create
    service = resource_def.service_factory(ctx.session)
    scope = ScopeContext(**payload.scope) if payload.scope else None
    options = WriteOptions(user_id=payload.user_id, scope=scope)

    valid_rows, errors = await _validate_rows(raw_rows, schema, ctx)

    imported_count = 0
    if not payload.dry_run:
        imported_count = await _persist_rows(
            service, valid_rows, payload, scope, options, ctx.session, ctx, errors
        )

    total_err_count = len(errors)
    errors_file_key: str | None = None
    truncated = total_err_count > MAX_EMBEDDED_ERRORS
    if truncated:
        errors_file_key = f"{IMPORT_STORAGE_PREFIX}/{ctx.job_id}_errors.json"
        full_err_bytes = json.dumps([e.model_dump() for e in errors], indent=2).encode(
            "utf-8"
        )
        try:
            await storage.upload(
                errors_file_key, full_err_bytes, content_type="application/json"
            )
        except Exception:
            errors_file_key = None
        errors = errors[:MAX_EMBEDDED_ERRORS]

    failed_count = (
        total_rows - imported_count
        if not payload.dry_run
        else len({e.row for e in errors})
    )
    result = ImportResult(
        total_rows=total_rows,
        imported_rows=imported_count,
        failed_rows=failed_count,
        errors=errors,
        total_errors=total_err_count,
        truncated=truncated,
        errors_file_key=errors_file_key,
        mode=payload.mode,
        dry_run=payload.dry_run,
    )
    await ctx.update_progress(
        100,
        f"Import complete: {imported_count} imported, {result.failed_rows} failed.",
    )
    return result.model_dump(mode="json")

import io

import openpyxl
from pydantic import BaseModel, Field

from fastapi_plantilla.core.crud.importer import (
    generate_import_template,
    parse_import_file,
)
from fastapi_plantilla.core.crud.schema import (
    ImportFormat,
    ImportMode,
    ImportResult,
    RowError,
)


class DummyItemCreate(BaseModel):
    """Test schema for item creation."""

    name: str = Field(..., description="Item Name")
    code: str = Field(..., description="Unique Code")
    price: float = Field(default=0.0, description="Price in EUR")
    description: str | None = None


def test_generate_template_csv() -> None:
    """Verify CSV template generation produces clean headers without double BOM."""
    content, media_type, filename = generate_import_template(
        DummyItemCreate, ImportFormat.CSV, "items"
    )
    assert media_type == "text/csv; charset=utf-8"
    assert filename == "items_template.csv"
    text = content.decode("utf-8-sig")
    assert not text.startswith("\ufeff")
    assert text.strip() == "name,code,price,description"


def test_generate_template_excel() -> None:
    """Verify Excel template generation produces exactly one header row."""
    content, media_type, filename = generate_import_template(
        DummyItemCreate, ImportFormat.EXCEL, "items"
    )
    expected_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert media_type == expected_type
    assert filename == "items_template.xlsx"
    assert len(content) > 100

    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    assert ws.max_row == 1
    headers = [cell.value for cell in ws[1]]
    assert headers == ["name", "code", "price", "description"]


def test_parse_csv_file_comma() -> None:
    """Verify parsing standard comma-delimited CSV."""
    csv_data = (
        b"name,code,price,description\n"
        b"Widget A,W1,19.99,First item\n"
        b"Widget B,W2,29.99,\n"
    )
    rows = parse_import_file(csv_data, "items.csv")
    assert len(rows) == 2
    assert rows[0]["name"] == "Widget A"
    assert rows[0]["code"] == "W1"
    assert rows[0]["price"] == "19.99"
    assert rows[1]["name"] == "Widget B"


def test_parse_csv_file_semicolon_and_bom() -> None:
    """Verify parsing semicolon-delimited CSV with UTF-8 BOM."""
    csv_data = "\ufeffname;code;price\nWidget C;W3;9.50\n".encode("utf-8")
    rows = parse_import_file(csv_data, "items.csv")
    assert len(rows) == 1
    assert rows[0]["name"] == "Widget C"
    assert rows[0]["code"] == "W3"


def test_parse_csv_ignores_empty_rows() -> None:
    """Verify blank lines are omitted from parsed results."""
    csv_data = b"name,code\nWidget X,WX\n\n   \nWidget Y,WY\n"
    rows = parse_import_file(csv_data, "items.csv")
    assert len(rows) == 2


def test_import_result_schema() -> None:
    """Verify ImportResult model attributes and counts."""
    errors = [RowError(row=1, field="code", message="Required", value=None)]
    res = ImportResult(
        total_rows=10,
        imported_rows=9,
        failed_rows=1,
        errors=errors,
        total_errors=1,
        truncated=False,
        mode=ImportMode.PARTIAL,
        dry_run=False,
    )
    assert res.total_rows == 10
    assert res.imported_rows == 9
    assert res.failed_rows == 1
    assert len(res.errors) == 1
    assert res.errors[0].field == "code"

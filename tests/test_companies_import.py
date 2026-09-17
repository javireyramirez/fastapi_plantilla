import io
import uuid
from collections.abc import AsyncGenerator
from typing import Any

import openpyxl
import pytest
from fastapi import FastAPI, status
from httpx import ASGITransport, AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from fastapi_plantilla.core.crud.importer import handle_import_job
from fastapi_plantilla.core.crud.repository import BaseRepository
from fastapi_plantilla.core.crud.schema import ImportJobPayload
from fastapi_plantilla.core.database import get_db_session
from fastapi_plantilla.core.mixins import generate_uuid7
from fastapi_plantilla.modules.auth.dependencies import (
    get_current_active_superuser,
    get_current_user,
)
from fastapi_plantilla.modules.auth.models import User
from fastapi_plantilla.modules.auth.schema import UserResponse
from fastapi_plantilla.modules.companies.models import Company
from fastapi_plantilla.modules.companies.routes import router as companies_router
from fastapi_plantilla.modules.jobs.exceptions import JobError
from fastapi_plantilla.modules.jobs.schema import JobContext
from fastapi_plantilla.modules.rbac.models import (
    RbacActions,
    Role,
    RoleAssignment,
    RolePermission,
    ScopeType,
    SystemModule,
)


class CompaniesAuthContext:
    """Authentication holder for switching users in test client."""

    def __init__(self, user: UserResponse) -> None:
        self.user = user


def _user_to_response(user: User) -> UserResponse:
    """Convert User ORM instance to UserResponse schema."""
    return UserResponse(
        id=user.id,
        name=user.name,
        email=user.email,
        email_verified=user.email_verified,
        is_active=user.is_active,
        is_system=user.is_system,
        is_super_admin=user.is_super_admin,
        created_at=user.created_at,
        updated_at=user.updated_at,
    )


@pytest.fixture
async def setup_companies_context(
    dbsession: AsyncSession,
) -> tuple[UserResponse, UserResponse]:
    """Create system modules, superadmin user, and regular member user."""
    mod_repo = BaseRepository(SystemModule, dbsession)
    for code, name in (
        ("companies", "Companies Module"),
        ("users", "Users Module"),
    ):
        if not await mod_repo.find_first(SystemModule.code == code):
            await mod_repo.create(
                {
                    "id": generate_uuid7(),
                    "code": code,
                    "name": name,
                    "description": f"{name} description",
                }
            )

    user_repo = BaseRepository(User, dbsession)
    admin_model = await user_repo.create(
        {
            "id": generate_uuid7(),
            "name": "Super Admin",
            "email": f"admin_import_{uuid.uuid4().hex[:6]}@example.com",
            "is_super_admin": True,
            "is_active": True,
        }
    )
    regular_model = await user_repo.create(
        {
            "id": generate_uuid7(),
            "name": "Regular User",
            "email": f"user_import_{uuid.uuid4().hex[:6]}@example.com",
            "is_super_admin": False,
            "is_active": True,
        }
    )

    role_repo = BaseRepository(Role, dbsession)
    perm_repo = BaseRepository(RolePermission, dbsession)
    assign_repo = BaseRepository(RoleAssignment, dbsession)

    role = await role_repo.create(
        {
            "id": generate_uuid7(),
            "name": f"Company Regular Role {uuid.uuid4().hex[:4]}",
            "slug": f"comp-regular-{uuid.uuid4().hex[:6]}",
        }
    )

    comp_mod = await mod_repo.find_first(SystemModule.code == "companies")
    assert comp_mod is not None

    for action in (RbacActions.READ,):
        await perm_repo.create(
            {
                "id": generate_uuid7(),
                "role_id": role.id,
                "module_id": comp_mod.id,
                "action": action,
                "scope": ScopeType.GLOBAL,
            }
        )

    await assign_repo.create(
        {
            "id": generate_uuid7(),
            "role_id": role.id,
            "entity_type": "USER",
            "entity_id": regular_model.id,
        }
    )

    await dbsession.commit()
    return _user_to_response(admin_model), _user_to_response(regular_model)


@pytest.fixture
async def companies_client(
    setup_companies_context: tuple[UserResponse, UserResponse],
    dbsession: AsyncSession,
) -> AsyncGenerator[tuple[AsyncClient, CompaniesAuthContext], None]:
    """Test client wired to companies router with switchable auth context."""
    admin_user, _ = setup_companies_context
    context = CompaniesAuthContext(admin_user)

    app = FastAPI()
    app.include_router(companies_router, prefix="/api")

    app.dependency_overrides[get_db_session] = lambda: dbsession
    app.dependency_overrides[get_current_user] = lambda: context.user
    app.dependency_overrides[get_current_active_superuser] = lambda: context.user

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as ac:
        yield ac, context


async def _run_test_job(
    job_id: uuid.UUID,
    payload: dict[str, Any],
    dbsession: AsyncSession,
) -> dict[str, Any]:
    """Simulate execution of the imports.validate background job handler."""

    async def dummy_update_progress(
        jid: uuid.UUID, token: int, progress: int, msg: str | None
    ) -> None:
        pass

    async def dummy_check_cancelled(jid: uuid.UUID, token: int) -> bool:
        return False

    ctx = JobContext(
        job_id=job_id,
        name="imports.validate",
        payload=ImportJobPayload(**payload),
        entity_type="companies",
        entity_id=None,
        lease_token=1,
        session=dbsession,
        _update_progress_fn=dummy_update_progress,
        _check_cancelled_fn=dummy_check_cancelled,
    )
    res = await handle_import_job(ctx)
    assert res is not None
    return res


@pytest.mark.anyio
async def test_download_import_template_excel(
    companies_client: tuple[AsyncClient, CompaniesAuthContext],
) -> None:
    """Verify Excel template download includes only one single header row."""
    client, _ = companies_client
    resp = await client.get("/api/companies/import-template?format=excel")
    assert resp.status_code == status.HTTP_200_OK
    assert (
        "application/vnd.openxmlformats-officedocument" in resp.headers["content-type"]
    )
    assert 'filename="companies_template.xlsx"' in resp.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)
    ws = wb.active
    assert ws.max_row == 1
    headers = [cell.value for cell in ws[1]]
    assert "name" in headers
    assert "nif" in headers
    assert "sector" in headers
    assert "id" not in headers
    assert "created_at" not in headers


@pytest.mark.anyio
async def test_download_import_template_csv(
    companies_client: tuple[AsyncClient, CompaniesAuthContext],
) -> None:
    """Verify CSV template download produces text/csv with schema headers."""
    client, _ = companies_client
    resp = await client.get("/api/companies/import-template?format=csv")
    assert resp.status_code == status.HTTP_200_OK
    assert "text/csv" in resp.headers["content-type"]
    assert 'filename="companies_template.csv"' in resp.headers["content-disposition"]
    text = resp.content.decode("utf-8-sig")
    assert not text.startswith("\ufeff")
    assert "name,nif,sector,website,description" in text


@pytest.mark.anyio
async def test_import_forbidden_for_unauthorized_user(
    companies_client: tuple[AsyncClient, CompaniesAuthContext],
    setup_companies_context: tuple[UserResponse, UserResponse],
) -> None:
    """Verify endpoint rejects requests without IMPORT permission."""
    client, context = companies_client
    _, regular_user = setup_companies_context
    context.user = regular_user

    resp = await client.get("/api/companies/import-template")
    assert resp.status_code == status.HTTP_403_FORBIDDEN


@pytest.mark.anyio
async def test_import_invalid_file_extension(
    companies_client: tuple[AsyncClient, CompaniesAuthContext],
) -> None:
    """Verify upload rejects unsupported file extensions like .pdf."""
    client, _ = companies_client
    files = {"file": ("test.pdf", b"dummy content", "application/pdf")}
    resp = await client.post("/api/companies/import", files=files)
    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    assert "Only .csv and .xlsx files are supported" in resp.json()["detail"]


@pytest.mark.anyio
async def test_import_rejects_xls_file(
    companies_client: tuple[AsyncClient, CompaniesAuthContext],
) -> None:
    """Verify upload rejects legacy .xls format explicitly."""
    client, _ = companies_client
    files = {"file": ("legacy.xls", b"dummy content", "application/vnd.ms-excel")}
    resp = await client.post("/api/companies/import", files=files)
    assert resp.status_code == status.HTTP_400_BAD_REQUEST
    assert "Only .csv and .xlsx files are supported" in resp.json()["detail"]


@pytest.mark.anyio
async def test_import_empty_file_fails(
    companies_client: tuple[AsyncClient, CompaniesAuthContext],
    dbsession: AsyncSession,
) -> None:
    """Verify that uploading an empty CSV file raises a JobError."""
    client, _ = companies_client
    csv_content = b"name,nif\n"  # Only header, no data rows
    files = {"file": ("companies.csv", csv_content, "text/csv")}
    resp = await client.post("/api/companies/import", files=files)
    assert resp.status_code == status.HTTP_202_ACCEPTED
    job_data = resp.json()
    job_id = uuid.UUID(job_data["id"])

    with pytest.raises(JobError, match="empty or contains no valid data rows"):
        await _run_test_job(job_id, job_data["payload"], dbsession)


@pytest.mark.anyio
async def test_import_atomic_success(
    companies_client: tuple[AsyncClient, CompaniesAuthContext],
    dbsession: AsyncSession,
) -> None:
    """Verify successful atomic batch import persists all records in DB."""
    client, _ = companies_client
    nif1 = f"A{uuid.uuid4().hex[:7]}".upper()
    nif2 = f"B{uuid.uuid4().hex[:7]}".upper()
    csv_content = (
        f"name,nif,sector\nAlpha Corp,{nif1},Tech\nBeta LLC,{nif2},Health\n"
    ).encode()

    files = {"file": ("companies.csv", csv_content, "text/csv")}
    data = {"mode": "atomic", "dry_run": "false"}
    resp = await client.post("/api/companies/import", files=files, data=data)
    assert resp.status_code == status.HTTP_202_ACCEPTED
    job_data = resp.json()
    job_id = uuid.UUID(job_data["id"])

    result = await _run_test_job(job_id, job_data["payload"], dbsession)
    assert result["imported_rows"] == 2
    assert result["failed_rows"] == 0
    assert len(result["errors"]) == 0

    stmt1 = select(Company).where(Company.nif == nif1)
    stmt2 = select(Company).where(Company.nif == nif2)
    c1 = (await dbsession.execute(stmt1)).scalar_one_or_none()
    c2 = (await dbsession.execute(stmt2)).scalar_one_or_none()
    assert c1 is not None
    assert c1.name == "Alpha Corp"
    assert c2 is not None
    assert c2.name == "Beta LLC"


@pytest.mark.anyio
async def test_import_atomic_rollback_on_validation_error(
    companies_client: tuple[AsyncClient, CompaniesAuthContext],
    dbsession: AsyncSession,
) -> None:
    """Verify ATOMIC mode rolls back entirely if any row fails validation."""
    client, _ = companies_client
    nif_valid = f"V{uuid.uuid4().hex[:7]}".upper()
    csv_content = (
        f"name,nif\nValid Corp,{nif_valid}\n,{uuid.uuid4().hex[:8]}\n"
    ).encode()

    files = {"file": ("companies.csv", csv_content, "text/csv")}
    data = {"mode": "atomic", "dry_run": "false"}
    resp = await client.post("/api/companies/import", files=files, data=data)
    assert resp.status_code == status.HTTP_202_ACCEPTED
    job_data = resp.json()
    job_id = uuid.UUID(job_data["id"])

    result = await _run_test_job(job_id, job_data["payload"], dbsession)
    assert result["imported_rows"] == 0
    assert result["failed_rows"] == 2
    assert len(result["errors"]) > 0

    stmt = select(Company).where(Company.nif == nif_valid)
    persisted = (await dbsession.execute(stmt)).scalar_one_or_none()
    assert persisted is None


@pytest.mark.anyio
async def test_import_atomic_rollback_on_db_nif_duplicate(
    companies_client: tuple[AsyncClient, CompaniesAuthContext],
    dbsession: AsyncSession,
) -> None:
    """Verify ATOMIC mode rolls back on database duplicate uniqueness constraint."""
    client, _ = companies_client
    existing_nif = f"D{uuid.uuid4().hex[:7]}".upper()

    existing_comp = Company(
        id=generate_uuid7(),
        name="Existing Company",
        nif=existing_nif,
        sector="Legal",
    )
    dbsession.add(existing_comp)
    await dbsession.commit()

    csv_content = f"name,nif\nNew Duplicate Corp,{existing_nif}\n".encode()
    files = {"file": ("companies.csv", csv_content, "text/csv")}
    data = {"mode": "atomic", "dry_run": "false"}
    resp = await client.post("/api/companies/import", files=files, data=data)
    assert resp.status_code == status.HTTP_202_ACCEPTED
    job_data = resp.json()
    job_id = uuid.UUID(job_data["id"])

    result = await _run_test_job(job_id, job_data["payload"], dbsession)
    assert result["imported_rows"] == 0
    assert result["failed_rows"] == 1
    assert any(err["field"] == "database" for err in result["errors"])


@pytest.mark.anyio
async def test_import_partial_mode(
    companies_client: tuple[AsyncClient, CompaniesAuthContext],
    dbsession: AsyncSession,
) -> None:
    """Verify PARTIAL mode inserts valid rows and logs failed rows."""
    client, _ = companies_client
    valid_nif = f"P{uuid.uuid4().hex[:7]}".upper()
    csv_content = f"name,nif\nPartial Valid,{valid_nif}\n,INV_PARTIAL_NIF\n".encode()

    files = {"file": ("companies.csv", csv_content, "text/csv")}
    data = {"mode": "partial", "dry_run": "false"}
    resp = await client.post("/api/companies/import", files=files, data=data)
    assert resp.status_code == status.HTTP_202_ACCEPTED
    job_data = resp.json()
    job_id = uuid.UUID(job_data["id"])

    result = await _run_test_job(job_id, job_data["payload"], dbsession)
    assert result["imported_rows"] == 1
    assert result["failed_rows"] == 1
    assert len(result["errors"]) >= 1

    stmt = select(Company).where(Company.nif == valid_nif)
    created = (await dbsession.execute(stmt)).scalar_one_or_none()
    assert created is not None
    assert created.name == "Partial Valid"


@pytest.mark.anyio
async def test_import_dry_run_mode(
    companies_client: tuple[AsyncClient, CompaniesAuthContext],
    dbsession: AsyncSession,
) -> None:
    """Verify dry_run simulation validates rows without creating database records."""
    client, _ = companies_client
    sim_nif = f"S{uuid.uuid4().hex[:7]}".upper()
    csv_content = f"name,nif\nSimulated Corp,{sim_nif}\n".encode()

    files = {"file": ("companies.csv", csv_content, "text/csv")}
    data = {"mode": "atomic", "dry_run": "true"}
    resp = await client.post("/api/companies/import", files=files, data=data)
    assert resp.status_code == status.HTTP_202_ACCEPTED
    job_data = resp.json()
    job_id = uuid.UUID(job_data["id"])

    result = await _run_test_job(job_id, job_data["payload"], dbsession)
    assert result["imported_rows"] == 0
    assert result["total_rows"] == 1
    assert result["dry_run"] is True

    stmt = select(Company).where(Company.nif == sim_nif)
    persisted = (await dbsession.execute(stmt)).scalar_one_or_none()
    assert persisted is None
